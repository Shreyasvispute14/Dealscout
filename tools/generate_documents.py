from __future__ import annotations

import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from reportlab.lib.colors import HexColor, black, white
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "frontend" / "public" / "downloads"
ASSETS = OUTPUT / "report-assets"
OUTPUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

NAVY = "#0B172A"
GREEN = "#10B981"
LIGHT = "#F1F5F9"
GREY = "#64748B"


def font(size: int, bold: bool = False):
    paths = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for path in paths:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def make_diagram(filename: str, title: str, boxes: list[tuple[str, int, int]], arrows: list[tuple[int, int]]) -> str:
    width, height = 1400, 700
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((20, 20, width - 20, height - 20), radius=24, fill="#F8FAFC", outline="#CBD5E1", width=3)
    draw.text((55, 45), title, fill=NAVY, font=font(34, True))
    positions: list[tuple[int, int, int, int]] = []
    for label, x, y in boxes:
        box = (x, y, x + 270, y + 115)
        positions.append(box)
        draw.rounded_rectangle(box, radius=18, fill="white", outline=GREEN, width=4)
        wrapped = textwrap.wrap(label, width=21)
        for line_no, line in enumerate(wrapped):
            draw.text((x + 24, y + 24 + line_no * 27), line, fill=NAVY, font=font(22, line_no == 0))
    for start, end in arrows:
        a, b = positions[start], positions[end]
        x1, y1 = (a[2], (a[1] + a[3]) // 2)
        x2, y2 = (b[0], (b[1] + b[3]) // 2)
        if x2 < x1:
            x1, y1 = ((a[0] + a[2]) // 2, a[3])
            x2, y2 = ((b[0] + b[2]) // 2, b[1])
        draw.line((x1, y1, x2, y2), fill="#475569", width=5)
        draw.polygon([(x2, y2), (x2 - 17, y2 - 10), (x2 - 17, y2 + 10)], fill="#475569")
    path = ASSETS / filename
    image.save(path, quality=95)
    return str(path)


def make_screen(filename: str, title: str, variant: str) -> str:
    image = Image.new("RGB", (1440, 820), "#F7F8F6")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1440, 76), fill="white")
    draw.rounded_rectangle((52, 18, 94, 60), radius=11, fill=GREEN)
    draw.text((65, 24), "D", fill=NAVY, font=font(22, True))
    draw.text((108, 24), "dealscout", fill=NAVY, font=font(25, True))
    draw.text((720, 27), "Compare     Insights     Resources", fill="#475569", font=font(17))
    if variant == "login":
        draw.rounded_rectangle((120, 130, 670, 740), radius=30, fill=NAVY)
        draw.text((180, 220), "Keep every comparison", fill="white", font=font(38, True))
        draw.text((180, 270), "within reach.", fill=GREEN, font=font(38, True))
        draw.rounded_rectangle((820, 150, 1270, 700), radius=24, fill="white", outline="#CBD5E1", width=2)
        draw.text((875, 220), "Welcome back", fill=NAVY, font=font(36, True))
        for y, label in [(315, "Email"), (425, "Password")]:
            draw.text((875, y), label, fill="#475569", font=font(18, True))
            draw.rounded_rectangle((875, y + 35, 1210, y + 92), radius=10, outline="#CBD5E1", width=2)
        draw.rounded_rectangle((875, 570, 1210, 630), radius=10, fill=NAVY)
        draw.text((1015, 586), "Sign in", fill="white", font=font(18, True))
    elif variant == "results":
        draw.rounded_rectangle((70, 115, 1370, 305), radius=30, fill=NAVY)
        draw.text((125, 155), "Your best options, side by side", fill="white", font=font(42, True))
        draw.rounded_rectangle((105, 260, 1335, 380), radius=20, fill="white", outline="#CBD5E1")
        draw.rounded_rectangle((70, 420, 1370, 565), radius=20, fill="#ECFDF5", outline=GREEN, width=3)
        draw.text((115, 455), "BEST PRICE", fill="#047857", font=font(17, True))
        draw.text((115, 495), "Apple iPhone 15 128GB", fill=NAVY, font=font(27, True))
        draw.text((1090, 475), "₹68,499", fill="#047857", font=font(36, True))
        for i, store in enumerate(["Flipkart", "Amazon", "Croma"]):
            y = 610 + i * 58
            draw.rectangle((90, y, 1350, y + 48), fill="white", outline="#E2E8F0")
            draw.text((120, y + 12), store, fill=NAVY, font=font(18, True))
            draw.text((600, y + 12), ["₹68,499", "₹69,999", "₹70,990"][i], fill=NAVY, font=font(18))
    elif variant == "insights":
        draw.text((70, 125), "Shopping insights", fill=NAVY, font=font(46, True))
        for i, (label, value) in enumerate([("COMPARISONS", "24"), ("UNIQUE SEARCHES", "11"), ("MOST SEARCHED", "iPhone 15")]):
            x = 70 + i * 435
            draw.rounded_rectangle((x, 205, x + 400, 365), radius=20, fill="white", outline="#CBD5E1")
            draw.text((x + 30, 245), label, fill=GREY, font=font(16, True))
            draw.text((x + 30, 290), value, fill=NAVY, font=font(32, True))
        draw.rounded_rectangle((70, 410, 920, 750), radius=20, fill="white", outline="#CBD5E1")
        draw.text((105, 450), "Recent comparisons", fill=NAVY, font=font(25, True))
        for i, text in enumerate(["iPhone 15      3 offers      ₹68,499", "MacBook Air     3 offers      ₹92,990", "Wireless Earbuds  2 offers      ₹1,299"]):
            draw.text((110, 525 + i * 65), text, fill="#475569", font=font(18))
        draw.rounded_rectangle((955, 410, 1370, 750), radius=20, fill="white", outline="#CBD5E1")
        draw.text((995, 450), "Store activity", fill=NAVY, font=font(25, True))
    else:
        draw.rounded_rectangle((60, 110, 1380, 430), radius=34, fill=NAVY)
        draw.text((125, 175), "More confidence.", fill="white", font=font(58, True))
        draw.text((125, 245), "Less searching.", fill=GREEN, font=font(58, True))
        draw.text((125, 335), "Compare prices, ratings, discounts and delivery details.", fill="#CBD5E1", font=font(22))
        draw.rounded_rectangle((115, 395, 1325, 570), radius=24, fill="white", outline="#CBD5E1", width=2)
        draw.text((160, 435), "Search for a product", fill=NAVY, font=font(22, True))
        draw.rounded_rectangle((160, 480, 995, 535), radius=10, outline="#CBD5E1", width=2)
        draw.text((185, 496), "iPhone 15", fill="#475569", font=font(17))
        draw.rounded_rectangle((1020, 480, 1280, 535), radius=10, fill=GREEN)
        draw.text((1070, 496), "Compare prices", fill="white", font=font(17, True))
        for i, text in enumerate(["Search once", "Compare clearly", "Spot the best"]):
            x = 110 + i * 440
            draw.rounded_rectangle((x, 635, x + 390, 760), radius=18, fill="white", outline="#CBD5E1")
            draw.text((x + 30, 675), text, fill=NAVY, font=font(23, True))
    draw.text((70, 785), title, fill="#94A3B8", font=font(14))
    path = ASSETS / filename
    image.save(path, quality=95)
    return str(path)


DIAGRAMS = {
    "er": make_diagram("er-diagram.png", "Entity Relationship Diagram", [("User", 70, 190), ("Search Log", 560, 190), ("Product Offer", 1050, 190), ("Store Source", 560, 450)], [(0, 1), (1, 2), (1, 3)]),
    "class": make_diagram("class-diagram.png", "Class Diagram", [("User\nid, name, email", 70, 190), ("SearchService\nsearch(), compare()", 560, 190), ("ProductOffer\nprice, rating, seller", 1050, 190), ("AnalyticsService\nsummarize()", 560, 450)], [(0, 1), (1, 2), (1, 3)]),
    "object": make_diagram("object-diagram.png", "Object Diagram", [("demoStudent : User", 70, 190), ("iphoneSearch : SearchLog", 560, 190), ("flipkartOffer : ProductOffer", 1050, 190), ("amazonOffer : ProductOffer", 560, 450)], [(0, 1), (1, 2), (1, 3)]),
    "component": make_diagram("component-diagram.png", "Component Diagram", [("React Web UI", 70, 190), ("FastAPI REST API", 560, 190), ("MongoDB", 1050, 190), ("Shopping Provider", 560, 450)], [(0, 1), (1, 2), (1, 3)]),
    "deployment": make_diagram("deployment-diagram.png", "Deployment Diagram", [("Browser", 70, 190), ("Frontend Container", 420, 190), ("Backend Container", 770, 190), ("MongoDB Container", 1120, 190)], [(0, 1), (1, 2), (2, 3)]),
    "usecase": make_diagram("use-case-diagram.png", "Use Case Diagram", [("Shopper", 70, 190), ("Search Product", 480, 120), ("Compare Offers", 900, 120), ("View Best Price", 900, 420), ("View Insights", 480, 420)], [(0, 1), (1, 2), (2, 3), (0, 4)]),
    "activity": make_diagram("activity-diagram.png", "Activity Diagram", [("Enter Query", 70, 190), ("Validate", 380, 190), ("Fetch Offers", 690, 190), ("Find Minimum", 1000, 190), ("Save & Display", 560, 450)], [(0, 1), (1, 2), (2, 3), (3, 4)]),
    "sequence": make_diagram("sequence-diagram.png", "Sequence Diagram", [("Shopper", 70, 190), ("Web UI", 390, 190), ("API", 710, 190), ("Database / Provider", 1030, 190), ("Comparison Result", 550, 450)], [(0, 1), (1, 2), (2, 3), (3, 4)]),
}

SCREENS = {
    "dashboard": make_screen("dashboard-screen.png", "Figure: DealScout dashboard and search input", "dashboard"),
    "login": make_screen("login-screen.png", "Figure: DealScout login screen", "login"),
    "results": make_screen("results-screen.png", "Figure: DealScout comparison result screen", "results"),
    "insights": make_screen("insights-screen.png", "Figure: DealScout shopping insights screen", "insights"),
}


def normal_page(title: str, purpose: str, points: list[str], explanation: str, image: str | None = None):
    return {"title": title, "purpose": purpose, "points": points, "explanation": explanation, "image": image}


PAGES = [
    {"type": "title", "title": "DEALSCOUT"},
    {"type": "certificate", "title": "CERTIFICATE"},
    normal_page("DECLARATION", "I/We hereby declare that the work presented in this report entitled “DealScout” is an original study completed for T.Y. B.Sc. (Computer Science), Semester V.", ["The implementation and documentation are prepared for academic evaluation.", "All external references are acknowledged in the bibliography.", "Student names and seat numbers are intentionally left editable."], "Placeholders should be completed and signed before final submission. The declaration confirms ownership of the submitted work while recognizing official documentation and software references."),
    normal_page("ACKNOWLEDGEMENT", "The completion of DealScout was made possible through guidance, review, testing, and access to academic resources.", ["Gratitude to the Principal and Department of Computer Science.", "Thanks to the Project Guide and Course Coordinator.", "Acknowledgement of internal and external examiners.", "Appreciation for classmates, family, and open-source communities."], "Replace this text with personal acknowledgements if required by the assigned teacher. Keep the tone formal, concise, and connected to the actual support received."),
    {"type": "index1", "title": "INDEX — PART I"},
    {"type": "index2", "title": "INDEX — PART II"},
    normal_page("ABSTRACT", "DealScout is a web application for comparing offers for the same product across multiple sellers. It combines search, normalization, recommendation, login, and search-history analysis in one responsive interface.", ["Accepts product names and category filters.", "Normalizes price, seller, rating, discount, and delivery fields.", "Highlights the lowest listed price.", "Stores searches for later insights.", "Supports a sample catalog and an optional live shopping provider."], "The solution reduces repeated manual searching. Its design separates the browser interface, REST API, data storage, and external provider adapter so that each part can be explained and maintained independently."),
    normal_page("CHAPTER 1 — PRELIMINARY INVESTIGATION", "Preliminary investigation establishes why DealScout is needed, what problem it solves, whether it is feasible, and what boundaries apply.", ["Study the existing manual comparison process.", "Identify measurable difficulties and user expectations.", "Define objectives before selecting technology.", "Assess technical, economic, operational, and schedule feasibility."], "The investigation found that buyers frequently move between tabs, copy prices manually, and overlook seller or delivery differences. A normalized comparison page offers a simpler decision path."),
    normal_page("1.1 PROBLEM IDENTIFICATION", "E-commerce platforms present similar products with different prices, discounts, ratings, seller names, and delivery estimates. Users must inspect each source separately.", ["Manual comparison consumes time.", "Different field formats create confusion.", "A low visible price may exclude context such as seller or delivery.", "Search history is normally scattered or unavailable.", "Users need a transparent explanation of the recommended deal."], "DealScout addresses these difficulties by displaying every normalized offer in one table and marking the minimum listed price. It does not claim that a preview price is live when no provider is connected."),
    normal_page("1.2 PROBLEM STATEMENT / DEFINITION", "To design and develop a responsive web system that searches product offers, converts them into a common structure, compares important values, identifies the lowest listed price, and stores search activity.", ["Input: product query, category, and source mode.", "Process: validation, retrieval, normalization, minimum-price selection, and logging.", "Output: best-deal card, comparison table, notices, and insights.", "Constraint: live results require a valid provider key."], "The problem definition is intentionally precise. It focuses on comparable offer data rather than promising purchase completion, payment handling, or guaranteed final checkout totals."),
    normal_page("1.3 PURPOSE, OBJECTIVES AND GOALS", "The purpose is to help users compare product offers with less effort and more context.", ["Provide a clean and accessible search form.", "Support mobile, laptop, and general-product categories.", "Display seller, price, rating, delivery, and discount.", "Calculate the lowest price using understandable logic.", "Protect insights behind a login session.", "Keep the application portable through containerized setup."], "Success is measured by whether a user can move from a query to a readable comparison and understand why one offer is highlighted. The interface must remain usable on desktop and mobile screens."),
    normal_page("1.4 TECHNICAL FEASIBILITY", "The required technologies are available, documented, and suitable for a student-scale application.", ["React and TypeScript render the responsive interface.", "FastAPI provides validated JSON endpoints.", "MongoDB stores search logs.", "Docker Compose packages frontend, backend, and database.", "An optional SerpApi adapter supplies live shopping results."], "The sample catalog guarantees that core functionality works without paid credentials. The application uses a typed HTTP boundary, clear data models, and server-side storage, making the architecture technically achievable."),
    normal_page("1.4 ECONOMIC FEASIBILITY", "Development can be completed with open-source tools and standard computer hardware.", ["No license fee for React, Python, FastAPI, or MongoDB Community.", "Docker Desktop is free for eligible educational use under its terms.", "The sample catalog has no API cost.", "Hosting and live-provider costs are optional and depend on usage."], "Economic feasibility is high because the MVP can run locally. The cost of internet access and a live provider is separated from the essential comparison workflow."),
    normal_page("1.4 OPERATIONAL AND SCHEDULE FEASIBILITY", "The application uses familiar search and table patterns, while the work can be divided into achievable stages.", ["Operational: clear labels, visible loading states, and readable errors.", "Operational: sign-in is needed only for private insights.", "Schedule: requirements and data model first.", "Schedule: backend, frontend, testing, and documentation follow.", "Schedule: teacher feedback is applied before submission."], "The modular sequence reduces project risk. Each stage produces a demonstrable output and supports independent correction without redesigning the entire system."),
    normal_page("1.5 PROJECT SCOPE", "DealScout covers product discovery and comparison for mobiles, laptops, and general items.", ["Search a predefined sample catalog.", "Optionally retrieve live shopping results.", "Compare normalized offers.", "Highlight the minimum price.", "Store and summarize search activity.", "Provide login, responsive design, downloadable resources, and one-click setup."], "The scope ends before checkout. DealScout is an information and decision-support application, not a retailer, payment processor, inventory authority, or delivery service."),
    normal_page("1.5 LIMITATIONS", "Every comparison system depends on the quality, freshness, and completeness of its source data.", ["Sample prices are not current retailer prices.", "Live mode requires a configured provider key.", "Taxes, exact variant, location, and shipping may affect final checkout cost.", "The demo login is not production identity management.", "The current catalog contains a limited number of products.", "Retailer links may be unavailable in sample mode."], "These limitations are shown honestly in the interface. Future work can add verified retailer partnerships, currency normalization, variant matching, and production authentication."),
    normal_page("CHAPTER 2 — REQUIREMENT SPECIFICATION", "Requirement specification translates the identified problem into hardware, software, technical, functional, data, performance, and security needs.", ["Requirements must be testable.", "Essential and optional capabilities are separated.", "Data fields match the actual API models.", "Security requirements reflect the current session design."], "This chapter provides a baseline for implementation and evaluation. A requirement should describe observable behavior rather than an internal preference."),
    normal_page("2.1 HARDWARE REQUIREMENTS", "DealScout is lightweight enough for an ordinary student computer.", ["Processor: dual-core 64-bit processor or better.", "Memory: minimum 4 GB RAM; 8 GB recommended for Docker.", "Storage: at least 3 GB free for images, dependencies, and database volume.", "Display: 1280×720 or higher recommended.", "Network: optional for sample mode, required for dependency download and live provider."], "Docker resource settings may be adjusted for older machines. Once images are built, the sample catalog and database can run on a local network without external API access."),
    normal_page("2.1 SOFTWARE REQUIREMENTS", "The portable setup uses Docker to avoid separate runtime installation.", ["Operating system: Windows 10/11, macOS, or modern Linux.", "Docker Desktop or Docker Engine with Compose.", "Modern browser: Chrome, Edge, Firefox, or Safari.", "Development stack: Node.js 24, Python 3.12, Nginx, and MongoDB 7 inside containers.", "Code editor is optional for normal use."], "The setup scripts call Docker Compose, build the services, start them in the background, and expose DealScout at localhost port 3000."),
    normal_page("2.2 FRONTEND TECHNICAL REQUIREMENTS", "The frontend must be responsive, type-safe, accessible, and independent of hard-coded backend addresses.", ["React components for pages and reusable interface blocks.", "TypeScript interfaces mirror API response models.", "TanStack Query manages server data and mutations.", "Tailwind CSS and shadcn components provide consistent styling.", "Relative /api paths work locally and behind the Nginx proxy."], "The interface renders its shell even when server data is unavailable. Interactive controls include clear data-testid values to support browser verification."),
    normal_page("2.2 BACKEND TECHNICAL REQUIREMENTS", "The backend must validate requests, use asynchronous database calls, and expose every route under /api.", ["FastAPI APIRouter organizes authentication, products, and analytics.", "Pydantic models validate requests and responses.", "Motor provides asynchronous MongoDB access.", "HTTPX connects to the optional shopping provider.", "Environment variables store connection settings and provider secrets."], "The server never exposes the provider key to browser code. Route modules share one database handle and return JSON structures that match frontend interfaces."),
    normal_page("2.2 DATABASE AND DEPLOYMENT REQUIREMENTS", "Search activity must survive service restarts and remain isolated from application source files.", ["MongoDB volume stores search logs.", "Indexes support recent-search sorting.", "Docker Compose starts MongoDB before application use.", "Nginx serves the frontend and proxies /api requests.", "The setup supports an optional SERPAPI_KEY environment variable."], "Although the college guideline describes relational tables, the implementation uses equivalent document collections. The report identifies entities, fields, keys, and relationships clearly."),
    normal_page("2.3 FUNCTIONAL REQUIREMENTS — ACCESS", "Access requirements define how visitors and signed-in users move through the application.", ["FR-01: Visitor can open the comparison page.", "FR-02: Visitor can open the sign-in page.", "FR-03: Valid credentials create an HTTP-only session cookie.", "FR-04: Invalid credentials return a clear error.", "FR-05: Signed-in user can view shopping insights.", "FR-06: Guest cannot retrieve protected analytics data."], "Authentication protects personal search summaries while leaving the core comparison experience publicly accessible."),
    normal_page("2.3 FUNCTIONAL REQUIREMENTS — COMPARISON", "The central functional requirements describe the complete search-to-recommendation flow.", ["FR-07: Accept a query between 2 and 80 characters.", "FR-08: Filter by All, Mobiles, Laptops, or General.", "FR-09: Select Sample prices or Live search.", "FR-10: Display normalized offers.", "FR-11: Identify the offer with minimum price.", "FR-12: Show a truthful notice when live data is unavailable.", "FR-13: Record every completed search."], "The highlighted result is based on listed price only. The UI retains seller, delivery, rating, and discount so the user can make a broader judgment."),
    normal_page("2.4 DATA REQUIREMENTS", "Data requirements define the minimum information needed to compare offers and summarize searches.", ["User: id, name, email, and role.", "ProductOffer: id, product name, category, source, seller, price, original price, rating, delivery, discount, URL, mode, flags, and retrieval time.", "SearchLog: id, query, category, mode, result count, best price, sources, and search time.", "AnalyticsSummary: totals, top query, source frequency, and recent searches."], "String UUID values are used as application identifiers. Prices use numeric values and timestamps are generated on the server."),
    normal_page("2.4 PERFORMANCE REQUIREMENTS", "The application should feel responsive for a small educational workload.", ["Local sample searches should normally return within one second.", "Live provider requests use a finite timeout.", "Search responses are limited to a small result set.", "Recent analytics reads are capped at 100 records.", "The UI displays loading text during network work.", "Responsive layouts avoid unnecessary horizontal scrolling except wide tables."], "Performance depends on hardware and provider availability. The design avoids unbounded database reads and does not block startup for index creation."),
    normal_page("2.4 SECURITY REQUIREMENTS", "Security requirements protect credentials, sessions, provider keys, and server data.", ["Provider secrets stay in backend environment variables.", "Sessions use HTTP-only cookies.", "Protected analytics validates the active session.", "Request bodies and query lengths are validated.", "Frontend never stores authentication tokens.", "Production deployments should use HTTPS, persistent sessions, password hashing, and restricted CORS."], "The bundled login is explicitly a preview account. It demonstrates access control but must be replaced before handling real user identities."),
    normal_page("CHAPTER 3 — DATABASE DESIGN AND END USERS", "Database design identifies the people who use DealScout and the entities required by their actions.", ["Visitor searches and compares offers.", "Signed-in shopper reviews insights.", "Future administrator manages catalog and provider settings.", "Evaluator or maintainer reviews behavior and records."], "User roles are deliberately simple. Permissions are applied at the analytics endpoint, while comparison remains available without registration."),
    normal_page("3.2 ENTITIES AND ATTRIBUTES", "The data model contains User, ProductOffer, SearchLog, and AnalyticsSummary entities.", ["User primary key: id.", "ProductOffer primary key: id.", "SearchLog primary key: id.", "SearchLog sources represent stores included in a comparison.", "AnalyticsSummary is calculated from SearchLog records rather than stored separately."], "Entity attributes correspond directly to validated backend models and handwritten frontend interfaces, reducing ambiguity across the HTTP boundary."),
    normal_page("3.2 ER DIAGRAM", "The ER diagram shows the logical relationship between a shopper, searches, offers, and store sources.", ["A user can review many search logs.", "A search can return many product offers.", "An offer belongs to one source at retrieval time.", "Analytics aggregates many search logs."], "The current preview does not persist a direct user foreign key on search logs; adding owner_id is recommended for a true multi-user production version.", DIAGRAMS["er"]),
    normal_page("3.3 USER TABLE / COLLECTION", "The User model represents an authenticated identity used to open protected insights.", ["id — string, primary application identifier.", "name — string, display name.", "email — string, login identifier.", "role — string, access role."], "The preview user is configured in server memory. A production implementation should use a users collection with unique email index, password hash, created time, and account status."),
    normal_page("3.3 PRODUCT OFFER STRUCTURE", "ProductOffer is the normalized structure used for every seller result.", ["price and original_price are numeric.", "currency defaults to INR.", "rating, delivery, discount, and URL may be absent.", "mode identifies sample or live retrieval.", "is_live and is_mocked prevent misleading labels.", "retrieved_at communicates freshness."], "Normalization allows a single comparison table and minimum-price function to work across different source schemas."),
    normal_page("3.3 SEARCH LOG STRUCTURE", "SearchLog records completed comparisons for later analysis.", ["id — unique string.", "query — submitted product text.", "category — selected filter.", "mode — demo or live.", "result_count — number of normalized offers.", "best_price — minimum price or null.", "searched_at — UTC server timestamp.", "sources — source names returned by the search."], "The searched_at descending index supports the recent-comparison view. In production, owner_id should be indexed with searched_at."),
    normal_page("3.3 RELATIONSHIPS AND DATA FLOW", "The browser submits a query, the backend produces ProductOffer objects, and one SearchLog is inserted before the response returns.", ["Logical User-to-SearchLog relationship: one-to-many.", "SearchLog-to-ProductOffer relationship: one-to-many during a response.", "Source-to-ProductOffer relationship: one-to-many.", "Analytics-to-SearchLog relationship: calculated aggregation."], "The application stores summary search data rather than full third-party payloads. This keeps the database compact and limits unnecessary retention."),
    normal_page("CHAPTER 4 — SYSTEM DESIGN", "System design explains the static structure, runtime objects, components, deployment nodes, user actions, workflows, and message sequences.", ["UML diagrams use one consistent model.", "Every diagram corresponds to implemented modules.", "External providers are shown as optional.", "The browser never contacts the provider directly."], "The diagrams that follow can be edited after teacher review while retaining the same architecture."),
    normal_page("4.1 CLASS DIAGRAM", "The class diagram connects data models to the services that use them.", ["User contains identity and role data.", "ProductOffer contains comparable fields.", "SearchService retrieves, normalizes, and compares.", "AnalyticsService summarizes logs.", "Authentication functions create and validate sessions."], "The diagram separates responsibilities so that UI rendering, business logic, and storage remain understandable.", DIAGRAMS["class"]),
    normal_page("4.2 OBJECT DIAGRAM", "The object diagram illustrates a single iPhone comparison at a moment in time.", ["demoStudent is an instance of User.", "iphoneSearch is an instance of SearchLog.", "flipkartOffer and amazonOffer are ProductOffer instances.", "The search object references returned source values."], "Object diagrams help explain how class definitions become actual runtime values.", DIAGRAMS["object"]),
    normal_page("4.3 COMPONENT DIAGRAM", "The component diagram displays the major deployable or logical software modules.", ["React web UI handles presentation and interaction.", "FastAPI REST API validates and coordinates work.", "MongoDB persists search logs.", "Shopping provider optionally supplies live offers.", "Nginx serves static files and proxies /api."], "All browser-server communication uses JSON over relative /api routes.", DIAGRAMS["component"]),
    normal_page("4.4 DEPLOYMENT DIAGRAM", "Docker Compose runs DealScout as three cooperating services on one machine.", ["Frontend container exposes port 3000.", "Backend container listens internally on 8001.", "MongoDB container stores data in a named volume.", "The browser reaches only the frontend origin.", "The backend reaches MongoDB and the optional provider."], "The same layout can run on Windows, macOS, Linux, or a compatible server.", DIAGRAMS["deployment"]),
    normal_page("4.5 USE CASE DIAGRAM", "Use cases describe behavior from the shopper's point of view.", ["Search product.", "Choose category and data source.", "Compare normalized offers.", "View best price.", "Sign in.", "View comparison insights.", "Download support resources."], "Advanced administration and purchasing are outside the current scope.", DIAGRAMS["usecase"]),
    normal_page("4.6 ACTIVITY DIAGRAM", "The main activity begins with user input and ends with a displayed and recorded comparison.", ["Enter and validate query.", "Read sample catalog or request provider.", "Normalize accepted results.", "Find minimum price.", "Insert search log.", "Display notice, best-price card, and offer table."], "If no offers match, the UI displays a useful empty state instead of a blank page.", DIAGRAMS["activity"]),
    normal_page("4.7 SEQUENCE DIAGRAM", "The sequence diagram emphasizes message order between shopper, UI, API, database, and provider.", ["UI sends GET /api/products/search.", "API optionally calls the provider.", "API constructs ProductOffer models.", "API saves a SearchLog.", "API returns SearchResponse JSON.", "UI renders and later invalidates analytics cache."], "Login follows a similar request sequence but creates an HTTP-only cookie.", DIAGRAMS["sequence"]),
    normal_page("SYSTEM ARCHITECTURE AND ALGORITHM", "DealScout uses a layered client-server architecture and a simple minimum-price recommendation algorithm.", ["Presentation layer: responsive React pages.", "API layer: validation and routing.", "Business layer: matching, normalization, and comparison.", "Data layer: MongoDB search logs.", "Integration layer: optional shopping API adapter."], "Algorithm: collect matching offers; discard entries without a usable price; compare numeric price values; select min(results, key=price); return its id as best_deal_id; render that row and card with emerald emphasis."),
    normal_page("CHAPTER 5 — INPUT AND OUTPUT SCREEN LAYOUT", "Screen layouts show how users enter information and how the system presents results.", ["Input screens include sign-in and search controls.", "Output screens include dashboard, offer comparison, and insights.", "Registration and add-record screens are documented as future scope.", "Each screen description identifies purpose, fields, actions, and validation."], "Actual implementation images are represented in the following figures and can be replaced with final deployment screenshots if requested by the teacher."),
    normal_page("5.1.1 LOGIN SCREEN", "The login screen grants access to private shopping insights.", ["Email input validates browser email format.", "Password input masks entered characters.", "Sign-in button submits credentials.", "Invalid credentials display a concise error.", "Preview account details support evaluation."], "The visual layout separates product value messaging from the focused sign-in card.", SCREENS["login"]),
    normal_page("5.1.2 REGISTRATION SCREEN", "Registration is not implemented in the current preview because only one evaluation account is required.", ["Proposed fields: full name, email, password, confirm password.", "Proposed validation: unique email and password policy.", "Proposed output: account creation confirmation and active session.", "Production passwords must be hashed, never stored directly."], "This documented future screen satisfies the planned multi-user extension without falsely claiming that registration is available today."),
    normal_page("5.1.3 ADD RECORD SCREEN", "Catalog editing is currently controlled in backend data rather than exposed to public users.", ["Proposed administrator fields: product name, category, store, seller, price, original price, rating, delivery, and URL.", "Required validation: positive price, recognized category, safe URL, and unique source offer.", "Actions: save, clear, cancel, and preview."], "An authenticated administrator role and audit log should be introduced before this screen is enabled."),
    normal_page("5.2.1 DASHBOARD", "The dashboard gives immediate access to brand navigation, mode selection, category filter, product query, examples, and feature explanation.", ["Hero communicates the service benefit.", "Search form is visually dominant.", "Sample and Live controls are clearly separated.", "Popular queries accelerate first use.", "Responsive cards explain the three-step flow."], "The page avoids academic language so it appears as a professional consumer product.", SCREENS["dashboard"]),
    normal_page("5.2.2 SEARCH RESULT SCREEN", "Search results place the best price before the detailed comparison table.", ["Query, result count, and source label remain visible.", "Preview-mode notice prevents misleading live claims.", "Best-price card identifies product, store, seller, delivery, price, and discount.", "Table supports horizontal scrolling on narrow screens.", "No-results state suggests shorter queries."], "Emerald is reserved for savings and the best-price recommendation.", SCREENS["results"]),
    normal_page("5.2.3 REPORT / INSIGHTS SCREEN", "The insights screen summarizes stored comparison activity for a signed-in user.", ["Comparison count.", "Unique search count.", "Most searched query.", "Recent comparison table.", "Store frequency list.", "Sign-in prompt for guests."], "The screen uses compact cards and tables rather than decorative charts so values remain easy to verify.", SCREENS["insights"]),
    normal_page("TESTING AND VALIDATION", "Testing confirms the main flow and important failure states across API, type system, and browser.", ["Public API root returns DealScout status.", "Valid search returns results and best_deal_id.", "One-character query returns HTTP 422.", "Valid login returns user and sets cookie.", "Protected analytics requires login.", "Frontend strict TypeScript check passes.", "Browser flow searches, signs in, and opens insights without console errors."], "A change is considered complete only when the exact failed check passes after correction."),
    normal_page("CONCLUSION AND FUTURE SCOPE", "DealScout demonstrates a complete comparison workflow with a professional responsive interface, validated API, data persistence, access control, downloadable resources, and portable setup.", ["Connect verified retailer or shopping-provider credentials.", "Add product-variant matching and currency normalization.", "Include shipping and tax in total-cost comparison.", "Replace preview login with persistent users and hashed passwords.", "Add price alerts, favorites, and historical trends.", "Create an administrator catalog editor."], "The current solution meets its defined objective while clearly separating implemented behavior from future enhancements."),
    normal_page("CHAPTER 6 — BIBLIOGRAPHY / REFERENCES", "The following sources support implementation and documentation.", ["FastAPI Documentation — https://fastapi.tiangolo.com/", "React Documentation — https://react.dev/", "MongoDB Documentation — https://www.mongodb.com/docs/", "Docker Documentation — https://docs.docker.com/", "MDN Web Docs — https://developer.mozilla.org/", "SerpApi Google Shopping API — https://serpapi.com/google-shopping-api", "SPPU and Kaveri College documentation instructions supplied for academic year 2026–27."], "All product names and retailer names are used for demonstration. Live data use must follow the chosen provider's terms, rate limits, attribution, and caching rules."),
    normal_page("APPENDIX — SUBMISSION AND VIVA CHECKLIST", "This final page helps prepare the report and system before review.", ["Fill student names, seat numbers, guide, coordinator, principal, examiner, date, and signatures.", "Ask the assigned teacher to review formatting and diagrams.", "Update index page numbers after final edits.", "Open the public or local app and test search, login, and insights.", "Explain sample data versus optional live data honestly.", "Carry the Word report, PDF report, viva guide, source code, and setup guide."], "Core viva explanation: DealScout receives a query, normalizes offers to one model, finds the minimum numeric price, stores a search log, and displays both the recommendation and supporting comparison rows."),
]

# The four chapter-openers below repeat the detailed first section that follows.
# Removing them keeps both generated formats at the required exact 50-page length.
PAGES = [page for page in PAGES if page["title"] not in {
    "CHAPTER 1 — PRELIMINARY INVESTIGATION",
    "CHAPTER 2 — REQUIREMENT SPECIFICATION",
    "CHAPTER 4 — SYSTEM DESIGN",
    "CHAPTER 5 — INPUT AND OUTPUT SCREEN LAYOUT",
}]

assert len(PAGES) == 50, f"Expected 50 pages, got {len(PAGES)}"


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_run(run, size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def add_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("DealScout  •  ")
    style_run(run, 9, color=GREY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_doc_title(document: Document, title: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(16)
    run = paragraph.add_run(title)
    style_run(run, 20, True, NAVY)


def add_doc_paragraph(document: Document, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    style_run(run, 12, bold)
    return paragraph


def add_doc_bullet(document: Document, text: str):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(4)
    style_run(paragraph.add_run(text), 12)


def add_title_page(document: Document):
    for _ in range(2): document.add_paragraph()
    for text, size, bold in [
        ("KANNADA SANGHA PUNE'S", 14, True),
        ("KAVERI COLLEGE OF ARTS, SCIENCE & COMMERCE, PUNE", 16, True),
        ("A PROJECT REPORT ON", 20, True),
        ("DEALSCOUT", 20, True),
        ("An E-Commerce Product Comparison and Best-Deal Recommendation System", 16, False),
        ("TYBSC(CS) (SEM – V)", 16, True),
        ("Academic Year 2026–27", 14, True),
    ]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(20 if size >= 20 else 12)
        style_run(paragraph.add_run(text), size, bold, NAVY)
    add_doc_paragraph(document, "Submitted by", True, WD_ALIGN_PARAGRAPH.CENTER)
    add_doc_paragraph(document, "Name 1: ______________________________    Seat No.: __________________", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_doc_paragraph(document, "Name 2: ______________________________    Seat No.: __________________", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_doc_paragraph(document, "In partial fulfilment of B. Sc. (Computer Science)", align=WD_ALIGN_PARAGRAPH.CENTER)


def add_certificate(document: Document):
    add_doc_title(document, "KAVERI COLLEGE OF ARTS, SCIENCE & COMMERCE")
    add_doc_paragraph(document, "PERMANENTLY AFFILIATED TO SAVITRIBAI PHULE PUNE UNIVERSITY, PUNE", True, WD_ALIGN_PARAGRAPH.CENTER)
    add_doc_paragraph(document, "Sr. No. 36, Shri G.M. Shetty Educational Complex, Ganeshnagar, Erandwane, Pune 411038.", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_doc_title(document, "CERTIFICATE")
    add_doc_paragraph(document, "This is to certify that Mr. / Ms. ______________________________, T.Y. BSc (Computer Science) Semester V, has successfully completed his/her Project Work titled “DealScout” in Computer Science Specialization as a part of CS-331-FP Project as prescribed by Savitribai Phule Pune University for academic year 2026–27.")
    add_doc_paragraph(document, "Seat No.: ____________________     Date: ____________________")
    table = document.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, role in enumerate(["Project Guide", "Course Coordinator", "Principal"]):
        set_cell_text(table.cell(0, index), role, True)
        set_cell_text(table.cell(1, index), "Name: __________________\nSignature: ______________")
    add_doc_paragraph(document, "Examined by", True, WD_ALIGN_PARAGRAPH.CENTER)
    exam = document.add_table(rows=2, cols=2)
    exam.style = "Table Grid"
    for index, role in enumerate(["Internal Examiner", "External Examiner"]):
        set_cell_text(exam.cell(0, index), role, True)
        set_cell_text(exam.cell(1, index), "Name: __________________\nSignature: ______________\nDate: __________________")


INDEX_ROWS = [
    ("1", "Preliminary Investigation — 1.1 Problem Identification; 1.2 Problem Statement; 1.3 Purpose, Objectives and Goals; 1.4 Feasibility Study; 1.5 Scope and Limitations", "8–15"),
    ("2", "Requirement Specification — 2.1 System Requirements; 2.2 Technical Requirement; 2.3 Functional Requirements; 2.4 Data Requirements", "16–25"),
    ("3", "Database Design — 3.1 End Users; 3.2 Entities, Attributes and ER Diagram; 3.3 Structures and Relationships", "26–32"),
    ("4", "System Design — Class, Object, Component, Deployment, Use Case, Activity and Sequence Diagrams", "33–40"),
    ("5", "Input and Output Screen Layout — Login, Registration, Add Record, Dashboard, Search Result and Insights", "41–46"),
    ("6", "Testing, Conclusion, Bibliography and Appendix", "47–50"),
]


def add_index(document: Document, part: int):
    add_doc_title(document, f"INDEX — PART {part}")
    rows = INDEX_ROWS[:3] if part == 1 else INDEX_ROWS[3:]
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, heading in enumerate(["Sr. No.", "Title", "Page No.", "Signature of Instructor"]):
        set_cell_text(table.cell(0, index), heading, True)
    for number, title, page in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], number)
        set_cell_text(cells[1], title)
        set_cell_text(cells[2], page)
        set_cell_text(cells[3], "")


def build_docx():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(24)
    add_footer(section)
    styles = document.styles
    if "Report Body" not in styles:
        body_style = styles.add_style("Report Body", WD_STYLE_TYPE.PARAGRAPH)
        body_style.font.name = "Times New Roman"
        body_style.font.size = Pt(12)
    for index, page in enumerate(PAGES):
        if index:
            document.add_page_break()
        page_type = page.get("type")
        if page_type == "title":
            add_title_page(document)
        elif page_type == "certificate":
            add_certificate(document)
        elif page_type == "index1":
            add_index(document, 1)
        elif page_type == "index2":
            add_index(document, 2)
        else:
            add_doc_title(document, page["title"])
            add_doc_paragraph(document, page["purpose"])
            for point in page["points"]:
                add_doc_bullet(document, point)
            add_doc_paragraph(document, page["explanation"])
            if page.get("image"):
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(page["image"], width=Inches(6.2))
                caption = document.add_paragraph(f"Figure {index + 1}: {page['title'].title()}")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_run(caption.runs[0], 10, False, GREY)
    document.core_properties.title = "DealScout Project Report — 50 Pages"
    document.core_properties.subject = "TYBSc Computer Science project documentation"
    document.core_properties.author = "DealScout"
    document.save(OUTPUT / "DealScout_Project_Report_50_Pages.docx")


def pdf_wrap(pdf, text: str, x: float, y: float, width_chars=96, leading=15, bullet=False):
    prefix = "• " if bullet else ""
    lines = textwrap.wrap(prefix + text, width=width_chars, subsequent_indent="  " if bullet else "") or [""]
    for line in lines:
        pdf.drawString(x, y, line)
        y -= leading
    return y


def pdf_header(pdf, title: str, page_no: int):
    width, height = A4
    pdf.setFillColor(HexColor(NAVY))
    pdf.rect(0, height - 28, width, 28, fill=1, stroke=0)
    pdf.setFillColor(HexColor(GREEN))
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(42, height - 18, "DEALSCOUT")
    pdf.setFillColor(HexColor(GREY))
    pdf.setFont("Helvetica", 9)
    pdf.drawCentredString(width / 2, 22, f"DealScout  •  Page {page_no} of 50")


def build_pdf():
    path = OUTPUT / "DealScout_Project_Report_50_Pages.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    for index, page in enumerate(PAGES, start=1):
        pdf_header(pdf, page["title"], index)
        page_type = page.get("type")
        if page_type == "title":
            pdf.setFillColor(HexColor(NAVY)); pdf.roundRect(45, 95, width - 90, height - 150, 24, fill=1, stroke=0)
            pdf.setFillColor(white); pdf.setFont("Times-Bold", 14); pdf.drawCentredString(width / 2, height - 120, "KANNADA SANGHA PUNE'S")
            pdf.setFont("Times-Bold", 17); pdf.drawCentredString(width / 2, height - 150, "KAVERI COLLEGE OF ARTS, SCIENCE & COMMERCE")
            pdf.setFillColor(HexColor(GREEN)); pdf.setFont("Times-Bold", 20); pdf.drawCentredString(width / 2, height - 265, "A PROJECT REPORT ON")
            pdf.setFont("Times-Bold", 36); pdf.drawCentredString(width / 2, height - 315, "DEALSCOUT")
            pdf.setFillColor(white); pdf.setFont("Times-Roman", 14); pdf.drawCentredString(width / 2, height - 350, "E-Commerce Product Comparison and Best-Deal Recommendation System")
            pdf.setFont("Times-Bold", 15); pdf.drawCentredString(width / 2, height - 405, "TYBSC(CS) (SEM – V)  •  Academic Year 2026–27")
            pdf.setFont("Times-Roman", 12); pdf.drawCentredString(width / 2, height - 485, "Name 1: __________________________   Seat No.: ______________")
            pdf.drawCentredString(width / 2, height - 515, "Name 2: __________________________   Seat No.: ______________")
            pdf.drawCentredString(width / 2, 135, "In partial fulfilment of B. Sc. (Computer Science)")
        elif page_type == "certificate":
            pdf.setFillColor(HexColor(NAVY)); pdf.setFont("Times-Bold", 15); pdf.drawCentredString(width / 2, height - 80, "KAVERI COLLEGE OF ARTS, SCIENCE & COMMERCE")
            pdf.setFont("Times-Roman", 9); pdf.drawCentredString(width / 2, height - 100, "PERMANENTLY AFFILIATED TO SAVITRIBAI PHULE PUNE UNIVERSITY, PUNE")
            pdf.setFont("Times-Bold", 20); pdf.drawCentredString(width / 2, height - 155, "CERTIFICATE")
            pdf.setFont("Times-Roman", 12); y = height - 205
            cert = "This is to certify that Mr. / Ms. ______________________________, T.Y. BSc (Computer Science) Semester V, has successfully completed his/her Project Work titled “DealScout” in Computer Science Specialization as a part of CS-331-FP Project as prescribed by Savitribai Phule Pune University for academic year 2026–27."
            y = pdf_wrap(pdf, cert, 65, y, 82, 18)
            y -= 22; pdf.drawString(65, y, "Seat No.: ____________________     Date: ____________________")
            y -= 110
            for x, role in [(65, "Project Guide"), (225, "Course Coordinator"), (405, "Principal")]:
                pdf.setFont("Times-Bold", 11); pdf.drawString(x, y, role); pdf.setFont("Times-Roman", 10); pdf.drawString(x, y - 28, "Name: ______________"); pdf.drawString(x, y - 48, "Signature: ___________")
            y -= 115; pdf.setFont("Times-Bold", 16); pdf.drawCentredString(width / 2, y, "Examined by")
            for x, role in [(100, "Internal Examiner"), (360, "External Examiner")]:
                pdf.setFont("Times-Bold", 11); pdf.drawString(x, y - 55, role); pdf.setFont("Times-Roman", 10); pdf.drawString(x, y - 80, "Name: __________________"); pdf.drawString(x, y - 100, "Signature: ______________"); pdf.drawString(x, y - 120, "Date: ___________________")
        elif page_type in {"index1", "index2"}:
            pdf.setFillColor(HexColor(NAVY)); pdf.setFont("Times-Bold", 20); pdf.drawCentredString(width / 2, height - 85, page["title"])
            rows = INDEX_ROWS[:3] if page_type == "index1" else INDEX_ROWS[3:]
            y = height - 135
            for number, title, page_range in rows:
                pdf.setStrokeColor(HexColor("#CBD5E1")); pdf.roundRect(48, y - 150, width - 96, 135, 8, fill=0, stroke=1)
                pdf.setFillColor(HexColor(GREEN)); pdf.setFont("Times-Bold", 15); pdf.drawString(65, y - 38, number)
                pdf.setFillColor(HexColor(NAVY)); pdf.setFont("Times-Roman", 11); pdf_wrap(pdf, title, 100, y - 35, 68, 15)
                pdf.setFont("Times-Bold", 10); pdf.drawRightString(width - 65, y - 38, page_range)
                pdf.setFont("Times-Roman", 9); pdf.drawRightString(width - 65, y - 125, "Instructor signature: __________________")
                y -= 165
        else:
            pdf.setFillColor(HexColor(NAVY)); pdf.setFont("Times-Bold", 20); pdf.drawCentredString(width / 2, height - 76, page["title"])
            pdf.setStrokeColor(HexColor(GREEN)); pdf.setLineWidth(3); pdf.line(55, height - 92, width - 55, height - 92)
            pdf.setFillColor(black); pdf.setFont("Times-Roman", 12); y = height - 125
            y = pdf_wrap(pdf, page["purpose"], 58, y, 88, 16)
            y -= 12
            for point in page["points"]:
                y = pdf_wrap(pdf, point, 72, y, 84, 15, True)
                y -= 3
            y -= 8
            y = pdf_wrap(pdf, page["explanation"], 58, y, 88, 16)
            if page.get("image") and y > 250:
                image_height = min(215, y - 65)
                pdf.drawImage(page["image"], 60, 55, width=width - 120, height=image_height, preserveAspectRatio=True, anchor="c")
        pdf.showPage()
    pdf.save()


VIVA_PAGES = [
    ("DealScout Viva Companion", "Use this guide to explain the system in simple, accurate language. Start with the problem, show the workflow, then discuss technology and limitations.", ["One-line definition: DealScout compares normalized product offers and highlights the lowest listed price.", "Demonstration order: search → comparison → best price → login → insights.", "Be honest that live mode needs a provider key."]),
    ("60-Second Introduction", "DealScout reduces the effort of visiting many stores separately. A user enters a product name, chooses a category and source mode, and receives normalized offers.", ["The backend validates the query.", "Offers use one common ProductOffer structure.", "The minimum numeric price becomes the best deal.", "Each search is stored for insights."]),
    ("Problem and Objectives", "The main problem is inconsistent and scattered e-commerce information.", ["Save comparison time.", "Improve clarity through a common table.", "Show seller, rating, discount, and delivery with price.", "Record searches for later review.", "Support portable setup."]),
    ("Architecture Answer", "The system follows a client-server architecture.", ["React and TypeScript run in the browser.", "Nginx serves the frontend and proxies /api.", "FastAPI handles validation and business logic.", "MongoDB stores search logs.", "An optional provider supplies live shopping results."]),
    ("Search Algorithm", "Explain the algorithm step by step without jargon.", ["Read and trim the query.", "Reject input shorter than two characters.", "Filter matching offers by query and category.", "Normalize every accepted source.", "Use minimum price to choose best_deal_id.", "Store the summary and return JSON."]),
    ("Database Answer", "The main stored entity is SearchLog; ProductOffer is the normalized response structure.", ["SearchLog: query, category, mode, count, best price, sources, timestamp.", "ProductOffer: product, source, seller, price, rating, delivery, discount.", "The recent-search query uses a descending timestamp index."]),
    ("Authentication Answer", "The preview login creates an HTTP-only cookie session.", ["The cookie travels automatically with same-origin requests.", "Browser JavaScript does not store a token.", "Analytics checks the session before returning data.", "Production would use persistent sessions and password hashing."]),
    ("API Questions", "REST endpoints exchange JSON between the browser and server.", ["GET /api/products/search — search and compare.", "POST /api/auth/login — create session.", "GET /api/auth/me — return current user or null.", "GET /api/analytics — protected search insights.", "HTTP 422 means request validation failed."]),
    ("Live vs Sample Data", "This distinction is important during the viva.", ["Sample catalog is local, stable, and always available.", "Live search requires SERPAPI_KEY on the backend.", "Without a key, preview results are explicitly labelled and not claimed as live.", "Provider keys never appear in frontend code."]),
    ("Testing Answer", "Testing covers the backend contract, strict frontend types, and the browser journey.", ["Positive API fields and status codes are asserted.", "A short invalid query proves validation.", "Typecheck catches interface mismatch.", "Browser test verifies search, login, insights, console, and network errors."]),
    ("Limitations and Future Scope", "A strong viva answer states limitations before enhancements.", ["Limited local catalog.", "No checkout or payment.", "Price may differ by tax, shipping, location, and variant.", "Preview login is not production identity.", "Future: verified providers, alerts, favorites, trends, and admin tools."]),
    ("Presentation Checklist", "Prepare the environment before entering the viva.", ["Open the app and keep the login ready.", "Search iPhone 15 in Sample prices.", "Show best-price logic and comparison rows.", "Sign in and open Insights.", "Keep Word/PDF report and source code available.", "Answer only what is implemented; call future work future work."]),
]


def build_short_pdf(filename: str, pages: list[tuple[str, str, list[str]]], subtitle: str):
    path = OUTPUT / filename
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    total = len(pages)
    for index, (title, intro, bullets) in enumerate(pages, start=1):
        pdf.setFillColor(HexColor(NAVY)); pdf.rect(0, height - 42, width, 42, fill=1, stroke=0)
        pdf.setFillColor(HexColor(GREEN)); pdf.setFont("Helvetica-Bold", 10); pdf.drawString(45, height - 26, "DEALSCOUT")
        pdf.setFillColor(HexColor(NAVY)); pdf.setFont("Times-Bold", 22); pdf.drawString(55, height - 95, title)
        pdf.setStrokeColor(HexColor(GREEN)); pdf.setLineWidth(3); pdf.line(55, height - 112, width - 55, height - 112)
        pdf.setFillColor(black); pdf.setFont("Times-Roman", 13); y = pdf_wrap(pdf, intro, 58, height - 150, 83, 19)
        y -= 25
        for bullet in bullets:
            y = pdf_wrap(pdf, bullet, 75, y, 78, 18, True)
            y -= 10
        pdf.setFillColor(HexColor(LIGHT)); pdf.roundRect(55, 85, width - 110, 90, 12, fill=1, stroke=0)
        pdf.setFillColor(HexColor(NAVY)); pdf.setFont("Helvetica-Bold", 11); pdf.drawString(75, 145, subtitle)
        pdf.setFont("Helvetica", 9); pdf_wrap(pdf, "Speak clearly, demonstrate the real flow, and separate implemented behavior from planned enhancements.", 75, 122, 82, 13)
        pdf.setFillColor(HexColor(GREY)); pdf.setFont("Helvetica", 9); pdf.drawCentredString(width / 2, 28, f"Page {index} of {total}")
        pdf.showPage()
    pdf.save()


SETUP_PAGES = [
    ("One-Click Setup", "DealScout includes Docker packaging for the frontend, backend, and database.", ["Install Docker Desktop.", "Extract the DealScout source folder.", "Windows: double-click setup.bat.", "macOS/Linux: run chmod +x setup.sh && ./setup.sh.", "Open http://localhost:3000."]),
    ("What Docker Starts", "Docker Compose creates a portable three-service environment.", ["frontend: builds React and serves it through Nginx.", "backend: runs FastAPI on the internal network.", "mongodb: stores search history in a named volume.", "Nginx forwards /api requests to the backend."]),
    ("Optional Live Provider", "Sample prices work without credentials. Live shopping requires a server-side key.", ["Create a .env file beside docker-compose.yml.", "Add SERPAPI_KEY=your_key.", "Run the setup script again.", "Never place the provider key in frontend code."]),
    ("Operations and Login", "Use these commands when maintaining the local installation.", ["Stop: docker compose down", "Start: docker compose up -d", "Logs: docker compose logs -f", "Reset data: docker compose down -v", "Preview login: preview@dealscout.local / dealscout123"]),
]


if __name__ == "__main__":
    build_docx()
    build_pdf()
    build_short_pdf("DealScout_Viva_Guide.pdf", VIVA_PAGES, "VIVA TIP")
    build_short_pdf("DealScout_Setup_Guide.pdf", SETUP_PAGES, "SETUP TIP")
    print("Generated:")
    for path in sorted(OUTPUT.glob("DealScout_*")):
        print(f"- {path.name}: {path.stat().st_size} bytes")